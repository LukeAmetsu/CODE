import math
import io
import base64
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from backend.database import db
from backend.calculators.beam_selector import find_lightest_beam

# Fixed pool of beams with predefined Lb (in ft)
FIXED_BEAMS = {
    "W8X10": 3, "W8X13": 3, "W8X15": 3, "W8X18": 5, "W8X21": 5,
    "W10X22": 5, "W10X26": 5,
    "W12X26": 5, "W12X35": 5, "W12X40": 7, "W12X45": 7, "W12X50": 6,
    "W14X30": 5, "W14X40": 7, "W14X48": 5, "W14X53": 5,
    "W16X57": 5,
    "W18X60": 6,
    "W21X62": 6, "W21X68": 6,
    "W24X68": 6
}

# --- AISC PIPE COLUMN DATABASE (A53 Gr B, Fy=35 ksi) ---
pipe_3_std = {
    8: 30.7, 9: 28.0, 10: 25.3, 11: 22.6, 12: 20.0, 13: 17.5,
    14: 15.1, 15: 13.1, 16: 11.6, 17: 10.2, 18: 9.13, 19: 8.19
}

pipe_35_std = {
    8: 40.3, 9: 37.6, 10: 34.8, 11: 31.9, 12: 29.0, 13: 26.2,
    14: 23.4, 15: 20.8, 16: 18.3, 17: 16.2, 18: 14.5, 19: 13.0, 20: 11.7, 21: 10.6, 22: 9.68
}

def get_beam_and_pipe(transverse_span, long_span, psf_load, pipe_height, Fy=50, omega=1.67):
    # The header beams span transversely from the building to the outer columns.
    # Therefore, the beam's span is the transverse_span (Row).
    # The spacing between these beams is the spacing between the outer columns (long_span / Column).
    # So the tributary width for the beam is the long_span.
    beam_span = transverse_span
    trib_width = long_span
    
    w_klf = (psf_load * trib_width) / 1000
    
    # 1. Select Beam using the centralized beam selector logic
    selector_inputs = {
        'design_method': 'ASD',
        'fy': Fy,
        'span_ft': beam_span,
        'w_load': w_klf,
        'restricted_pool': FIXED_BEAMS,
        'check_deflection': False  # Skip deflection check to match simple moment matrix behavior
    }
    
    res = find_lightest_beam(selector_inputs)
    
    beam_shape = "NG"
    beam_lb = 0.0
    
    if res.get('candidates') and len(res['candidates']) > 0:
        winner = res['candidates'][0]
        
        # Ensure correct formatting: W8X10 -> W8x10
        name = winner['name']
        parts = name.split('X')
        if len(parts) == 2:
            beam_shape = f"{parts[0]}x{parts[1]}"
        else:
            beam_shape = name
            
        # Use the fixed Lb from the restricted pool
        beam_lb = FIXED_BEAMS.get(name.upper(), 0)

    # 2. Check Pipe Leg Capacity
    # The leg supports one end of the header beam. Total load on beam = w_klf * beam_span.
    # The leg reaction is half of the total load.
    P_leg = (w_klf * beam_span) / 2
    kl = int(round(pipe_height))
    cap_3 = pipe_3_std.get(kl, 0)
    cap_35 = pipe_35_std.get(kl, 0)
    
    if P_leg <= cap_3:
        leg_check = "P3"
    elif P_leg <= cap_35:
        leg_check = "P3.5"
    else:
        leg_check = "NG"
        
    return {
        "beam": beam_shape,
        "leg": leg_check,
        "p_leg": round(P_leg, 1)
    }

def generate_matrix_data(inputs):
    load = float(inputs.get("load", 300))
    pipe_height = float(inputs.get("pipe_height", 14.0))
    
    transverse_spans = list(range(4, 31))
    longitudinal_spans = list(range(4, 32, 2))
    
    rows = []
    for t_span in transverse_spans:
        row_data = {"t_span": t_span}
        cells = []
        for l_span in longitudinal_spans:
            result = get_beam_and_pipe(t_span, l_span, load, pipe_height)
            cells.append({
                "l_span": l_span,
                "beam": result["beam"],
                "leg": result["leg"],
                "p_leg": result["p_leg"]
            })
        row_data["cells"] = cells
        rows.append(row_data)
        
    return {
        "load": load,
        "pipe_height": pipe_height,
        "long_spans": longitudinal_spans,
        "rows": rows,
        "fixed_beams": FIXED_BEAMS
    }

def generate_matrix(inputs):
    try:
        data = generate_matrix_data(inputs)
        return {"status": "success", "data": data}
    except Exception as e:
        import traceback
        return {"error": str(e), "trace": traceback.format_exc()}

def export_matrix_to_excel(inputs):
    try:
        data = generate_matrix_data(inputs)
        
        columns = ["REQUIRED SPAN"]
        for ls in data["long_spans"]:
            columns.append(f"{ls}")
            
        df_data = []
        style_map = [] 
        
        for r_idx, row in enumerate(data["rows"]):
            row_list = [row["t_span"]]
            row_styles = [None]
            for cell in row["cells"]:
                row_list.append(f"{cell['beam']}\n{cell['p_leg']}k")
                row_styles.append(cell["leg"])
                
            df_data.append(row_list)
            style_map.append(row_styles)
            
        df = pd.DataFrame(df_data, columns=columns)
        
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Shed Matrix')
            workbook = writer.book
            worksheet = writer.sheets['Shed Matrix']
            
            fill_p3 = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            fill_p35 = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
            fill_ng = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            
            center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            for col_idx, col in enumerate(worksheet.columns, 1):
                worksheet.column_dimensions[col[0].column_letter].width = 12
                for cell in col:
                    cell.alignment = center_align
            
            worksheet.column_dimensions['A'].width = 18
            
            for r_idx, styles in enumerate(style_map, 2):
                for c_idx, leg_status in enumerate(styles, 1):
                    if leg_status:
                        cell = worksheet.cell(row=r_idx, column=c_idx)
                        if leg_status == "P3":
                            cell.fill = fill_p3
                        elif leg_status == "P3.5":
                            cell.fill = fill_p35
                        elif leg_status == "NG":
                            cell.fill = fill_ng

            # Create Lb Reference Sheet
            df_lb = pd.DataFrame(list(FIXED_BEAMS.items()), columns=["Beam Size", "Lb (FT)"])
            df_lb.to_excel(writer, index=False, sheet_name='Lb Reference')
            ws_lb = writer.sheets['Lb Reference']
            ws_lb.column_dimensions['A'].width = 15
            ws_lb.column_dimensions['B'].width = 15
            for col in ws_lb.columns:
                for cell in col:
                    cell.alignment = center_align

        b64 = base64.b64encode(output.getvalue()).decode('utf-8')
        return {"status": "success", "filename": f"Shed_Matrix_{data['load']}psf_{data['pipe_height']}ft.xlsx", "b64data": b64}
        
    except Exception as e:
        import traceback
        return {"error": str(e), "trace": traceback.format_exc()}
