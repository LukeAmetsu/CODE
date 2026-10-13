import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional


def export_takeoff_to_excel(entries: List[Dict[str, Any]], title: str = "Base de Quantitativos de Fundações") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo Geral"

    # Header Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    data_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Title Block
    ws.merge_cells("A1:K1")
    ws["A1"] = f"CENTRAL DE ENGENHARIA - {title.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Categoria",
        "Projeto / Obra",
        "Equipamento / Estrutura",
        "Tipo de Fundação",
        "Nº Prancha / DWG",
        "Rev.",
        "fck (MPa)",
        "Concreto (m³)",
        "Aço CA-50 (kg)",
        "Taxa Aço (kg/m³)",
        "Fôrma (m²)",
        "Escavação (m³)"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 28

    current_row = 4
    for e in entries:
        cat_label = "Termelétrica (UTE)" if e.get('category') == 'termoeletrica' else "Linha Transmissão (LT)"
        vol = float(e.get('concrete_volume_m3') or 0.0)
        steel = float(e.get('steel_weight_kg') or 0.0)
        ratio = float(e.get('steel_ratio_kg_m3') or 0.0)
        form = float(e.get('formwork_area_m2') or 0.0)
        exc = float(e.get('excavation_m3') or 0.0)
        fck = float(e.get('fck_concrete') or 0.0)

        row_vals = [
            cat_label,
            e.get('project_name', ''),
            e.get('equipment_name', ''),
            e.get('structure_type', ''),
            e.get('drawing_code', ''),
            str(e.get('revision', '')),
            fck,
            vol,
            steel,
            ratio,
            form,
            exc
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00" if col_idx in [8, 10, 11, 12] else ("#,##0" if col_idx == 9 else "0.0")
            elif col_idx in [1, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Total Row
    tot_row = current_row
    ws.cell(row=tot_row, column=1, value="TOTAL GERAL").font = bold_font
    ws.cell(row=tot_row, column=1).alignment = Alignment(horizontal="center")
    for c in range(1, 8):
        ws.cell(row=tot_row, column=c).border = thin_border
    
    # Formulas for totals
    ws.cell(row=tot_row, column=8, value=f"=SUM(H4:H{tot_row-1})").number_format = "#,##0.00"
    ws.cell(row=tot_row, column=9, value=f"=SUM(I4:I{tot_row-1})").number_format = "#,##0"
    ws.cell(row=tot_row, column=10, value=f"=IF(H{tot_row}>0, I{tot_row}/H{tot_row}, 0)").number_format = "#,##0.00"
    ws.cell(row=tot_row, column=11, value=f"=SUM(K4:K{tot_row-1})").number_format = "#,##0.00"
    ws.cell(row=tot_row, column=12, value=f"=SUM(L4:L{tot_row-1})").number_format = "#,##0.00"

    for c in range(8, 13):
        ws.cell(row=tot_row, column=c).font = bold_font
        ws.cell(row=tot_row, column=c).border = thin_border
        ws.cell(row=tot_row, column=c).fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row == 1:
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
