import unittest
import os
import tempfile
from backend.calculators.comb_bases_mietc import (
    parse_mietc_workbook,
    calculate_single_base_combinations,
    batch_calculate_sheet_bases,
    export_mietc_results_to_excel,
    auto_detect_subsequent_pairs,
    calculate_combined_footing_envelopes
)

class TestMietcCombinations(unittest.TestCase):
    def setUp(self):
        self.ref_file = os.path.join(os.path.dirname(__file__), "..", "..", "Refs", "Cópia de Cargas nas bases MIETC.xlsx")
        self.has_ref = os.path.exists(self.ref_file)

    def test_parse_workbook(self):
        if not self.has_ref:
            self.skipTest("Reference Excel file not found.")
        
        parsed = parse_mietc_workbook(self.ref_file)
        self.assertIn("sheets", parsed)
        self.assertIn("sheet_names", parsed)
        self.assertGreater(len(parsed["sheet_names"]), 0)
        
        # Test Cargas 01
        sheet1 = parsed["sheets"].get("Cargas 01")
        self.assertIsNotNone(sheet1)
        self.assertIn("bases_data", sheet1)
        self.assertIn("BT1", sheet1["bases_data"])
        self.assertIn("PP", sheet1["bases_data"]["BT1"])
        self.assertIn("V", sheet1["bases_data"]["BT1"]["PP"])
        self.assertIn("X", sheet1["bases_data"]["BT1"]["PP"])
        self.assertIn("Y", sheet1["bases_data"]["BT1"]["PP"])

    def test_calculate_single_base(self):
        if not self.has_ref:
            self.skipTest("Reference Excel file not found.")

        parsed = parse_mietc_workbook(self.ref_file)
        sheet1 = parsed["sheets"]["Cargas 01"]
        base_bt1 = sheet1["bases_data"]["BT1"]

        # Run ELU combinations
        res = calculate_single_base_combinations("BT1", base_bt1, standard="NBR 8681", method="ELU_NORMAL")
        self.assertIn("combinations", res)
        self.assertIn("envelopes", res)
        self.assertGreater(len(res["combinations"]), 0)

        combos = res["combinations"]
        for c in combos:
            self.assertIn("V", c)
            self.assertIn("X", c)
            self.assertIn("Y", c)
            self.assertIn("H_res", c)
            self.assertGreaterEqual(c["H_res"], 0.0)

        envs = res["envelopes"]
        self.assertIn("V_max", envs)
        self.assertIn("V_min", envs)
        self.assertIn("H_max", envs)
        self.assertGreaterEqual(envs["V_max"]["value"], envs["V_min"]["value"])

    def test_batch_calculate_and_export_excel(self):
        if not self.has_ref:
            self.skipTest("Reference Excel file not found.")

        parsed = parse_mietc_workbook(self.ref_file)
        sheet1 = parsed["sheets"]["Cargas 01"]

        batch_res = batch_calculate_sheet_bases(sheet1, standard="NBR 8681", method="ELU_NORMAL")
        self.assertIn("summary_table", batch_res)
        self.assertIn("global_kpi", batch_res)
        self.assertEqual(len(batch_res["summary_table"]), len(sheet1["bases_data"]))

        # Check export
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp_path = tmp.name

        try:
            out_file = export_mietc_results_to_excel(batch_res, tmp_path)
            self.assertTrue(os.path.exists(out_file))
            self.assertGreater(os.path.getsize(out_file), 1000)

            # Verificar a aba executiva de resumo de cargas para avaliação de terceiros
            import openpyxl
            wb = openpyxl.load_workbook(out_file)
            self.assertIn("Resumo_Cargas_Consideradas", wb.sheetnames)
            self.assertIn("Resumo_Envoltórias", wb.sheetnames)
            self.assertIn("Todas_Combinações", wb.sheetnames)
            self.assertEqual(wb.sheetnames[0], "Resumo_Cargas_Consideradas")

            ws_cargas = wb["Resumo_Cargas_Consideradas"]
            self.assertIn("RESUMO DAS CARGAS CONSIDERADAS", str(ws_cargas.cell(1, 1).value))
            self.assertEqual(ws_cargas.cell(5, 2).value, "Ação (Sigla)")
            self.assertEqual(ws_cargas.cell(5, 7).value, "Pode Zerar (Alívio Crítico)")
            wb.close()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_exclusive_groups_and_requirements(self):
        """Verifica que ações com o mesmo grupo exclusivo (XOR) nunca atuam juntas e que AND é respeitado"""
        base_loads = {
            'PP': {'V': 100.0, 'X': 0.0, 'Y': 0.0},
            'SCM': {'V': 20.0, 'X': 0.0, 'Y': 0.0},
            'CVX+': {'V': 0.0, 'X': 30.0, 'Y': 0.0},
            'CVY+': {'V': 0.0, 'X': 0.0, 'Y': 30.0},
            'TCQ': {'V': 40.0, 'X': 5.0, 'Y': 0.0},
            'TCP': {'V': 15.0, 'X': 2.0, 'Y': 0.0},
            'SCA': {'V': 12.0, 'X': 0.0, 'Y': 0.0},
        }

        # Regras com XOR e AND
        options = {
            'action_rules': {
                'CVX+': {'group': 'Vento', 'requires': '', 'type': 'Vento (W)'},
                'CVY+': {'group': 'Vento', 'requires': '', 'type': 'Vento (W)'},
                'TCQ': {'group': 'Talha', 'requires': '', 'type': 'Ponte Rolante / Guindaste (Q)'},
                'TCP': {'group': 'Talha', 'requires': '', 'type': 'Ponte Rolante / Guindaste (Q)'},
                'SCA': {'group': '', 'requires': '', 'type': 'Sobrecarga Acidental (SCA)'},
            },
            'custom_loads': [
                {'name': 'Empuxo Solo Seco', 'type': 'Solo Seco (Empuxo)', 'value': 35.0, 'group': 'Solo', 'requires': '', 'direction': 'X'},
                {'name': 'Empuxo Solo Saturado', 'type': 'Solo Saturado (Empuxo)', 'value': 25.0, 'group': 'Solo', 'requires': 'Pressão da Água', 'direction': 'X'},
                {'name': 'Pressão da Água', 'type': 'Pressão Hidrostática / Água', 'value': 18.0, 'group': '', 'requires': '', 'direction': 'X'}
            ]
        }

        res = calculate_single_base_combinations("TEST_BASE", base_loads, standard="NBR 8681", method="ELU_NORMAL", options=options)
        combos = res["combinations"]
        self.assertGreater(len(combos), 0)

        for c in combos:
            factors = c.get("factors", {})
            # 1. Teste XOR Vento: CVX+ e CVY+ nunca podem ter fatores positivos simultaneamente
            self.assertFalse(
                factors.get('CVX+', 0) > 0 and factors.get('CVY+', 0) > 0,
                f"Combinação violou XOR Vento (CVX+ e CVY+ simultâneos): {c['title']}"
            )

            # 2. Teste XOR Talha: TCQ e TCP nunca podem ter fatores positivos simultaneamente
            self.assertFalse(
                factors.get('TCQ', 0) > 0 and factors.get('TCP', 0) > 0,
                f"Combinação violou XOR Talha (TCQ e TCP simultâneos): {c['title']}"
            )

            # 3. Teste XOR Solo: Solo Seco e Solo Saturado nunca podem ter fatores positivos simultaneamente
            self.assertFalse(
                factors.get('Empuxo Solo Seco', 0) > 0 and factors.get('Empuxo Solo Saturado', 0) > 0,
                f"Combinação violou XOR Solo: {c['title']}"
            )

            # 4. Teste AND: Empuxo Solo Saturado requer Pressão da Água
            if factors.get('Empuxo Solo Saturado', 0) > 0:
                self.assertGreater(
                    factors.get('Pressão da Água', 0), 0,
                    f"Combinação violou AND (Empuxo Solo Saturado sem Pressão da Água): {c['title']}"
                )

            # 5. Teste Condição Física: Solo Seco desativa Água
            if factors.get('Empuxo Solo Seco', 0) > 0:
                self.assertEqual(
                    factors.get('Pressão da Água', 0), 0,
                    f"Combinação violou condição física (Solo Seco com Pressão da Água): {c['title']}"
                )

    def test_formulas_no_trailing_plus_and_six_envelopes(self):
        """Verifica que as fórmulas geradas nunca terminam com '+' isolado e que as 6 envoltórias estão calculadas"""
        base_loads = {
            'PP': {'V': 150.0, 'X': 0.0, 'Y': 0.0},
            'CP': {'V': 50.0, 'X': 0.0, 'Y': 0.0},
            'SCM': {'V': 30.0, 'X': 0.0, 'Y': 0.0},
            'SCP': {'V': 25.0, 'X': 0.0, 'Y': 0.0},
            'T+': {'V': 5.0, 'X': 10.0, 'Y': 0.0},
            'CVX+': {'V': 0.0, 'X': 45.0, 'Y': 0.0},
            'CVY+': {'V': 0.0, 'X': 0.0, 'Y': 40.0},
        }

        res = calculate_single_base_combinations("B_TEST", base_loads, standard="NBR 8681", method="TODAS")
        combos = res["combinations"]
        self.assertGreater(len(combos), 0)

        for c in combos:
            formula = c.get("formula", "").strip()
            # A fórmula nunca deve terminar com o caractere '+' solto
            self.assertFalse(formula.endswith('+'), f"Fórmula terminou com '+': {formula}")
            # Ações com '+' no nome como CVX+, CVY+, T+ devem aparecer entre parênteses: (CVX+), (CVY+), (T+)
            for act in ['CVX+', 'CVY+', 'T+']:
                if act in formula:
                    self.assertIn(f"({act})", formula, f"Ação {act} não formatada com parênteses em: {formula}")

        # Verifica a estrutura completa das 6 envoltórias
        envs = res["envelopes"]
        self.assertIn("peaks", envs)
        self.assertIn("V_max", envs)
        self.assertIn("V_min", envs)
        self.assertIn("X_max", envs)
        self.assertIn("Y_max", envs)
        self.assertIn("H_max", envs)

        # 1. Picos Globais
        peaks = envs["peaks"]
        self.assertIn("V_max", peaks)
        self.assertIn("X_max", peaks)
        self.assertIn("Y_max", peaks)
        self.assertIn("H_max", peaks)

        # 2. V_max e concomitantes
        self.assertIn("concomitant_X", envs["V_max"])
        self.assertIn("concomitant_Y", envs["V_max"])
        self.assertIn("concomitant_H", envs["V_max"])

        # 3. V_min e concomitantes
        self.assertIn("concomitant_X", envs["V_min"])
        self.assertIn("concomitant_Y", envs["V_min"])
        self.assertIn("concomitant_H", envs["V_min"])

        # 4. X_max e concomitantes
        self.assertIn("concomitant_V", envs["X_max"])
        self.assertIn("concomitant_Y", envs["X_max"])
        self.assertIn("concomitant_H", envs["X_max"])

        # 5. Y_max e concomitantes
        self.assertIn("concomitant_V", envs["Y_max"])
        self.assertIn("concomitant_X", envs["Y_max"])
        self.assertIn("concomitant_H", envs["Y_max"])

        # 6. H_max e concomitantes
        self.assertIn("concomitant_V", envs["H_max"])
        self.assertIn("concomitant_X", envs["H_max"])
        self.assertIn("concomitant_Y", envs["H_max"])

        # Testar resumo de batch
        fake_sheet = {
            'sheet_name': 'Sheet1',
            'bases_data': {'B_TEST': base_loads}
        }
        batch = batch_calculate_sheet_bases(fake_sheet, standard="NBR 8681", method="ELU_NORMAL")
        row = batch["summary_table"][0]
        # Todos os 6 envelopes devem estar preenchidos no summary_table
        self.assertIn("peak_V", row)
        self.assertIn("peak_X", row)
        self.assertIn("peak_Y", row)
        self.assertIn("peak_H", row)
        self.assertIn("V_max", row)
        self.assertIn("V_min", row)
        self.assertIn("X_max", row)
        self.assertIn("Y_max", row)
        self.assertIn("H_max", row)
        self.assertIn("X_max_V", row)
        self.assertIn("Y_max_V", row)
        self.assertIn("H_max_V", row)

    def test_zeroed_load_conservative_uplift(self):
        """Verifica que cargas verticais zeráveis são desconsideradas para encontrar a máxima tração (arrancamento)"""
        base_loads = {
            'PP': {'V': 50.0, 'X': 0.0, 'Y': 0.0},
            'CVX+': {'V': -150.0, 'X': 60.0, 'Y': 0.0}, # Vento de tombamento causando tração (-150 kN)
            'SCA': {'V': 80.0, 'X': 0.0, 'Y': 0.0},     # Sobrecarga vertical (+80 kN) que aliviaria a tração se mantida
        }
        options = {
            'action_rules': {
                'PP': {'can_zero': False},
                'CVX+': {'type': 'Vento (W)', 'can_zero': False},
                'SCA': {'type': 'Sobrecarga Acidental (SCA)', 'can_zero': True} # Usuário indicou que SCA pode ser zerada
            },
            'enable_zero_load_scenarios': True
        }
        res = calculate_single_base_combinations("Sapata_Torre", base_loads, standard="NBR 8681", method="ELU_NORMAL", options=options)
        
        # O caso com SCA ativa daria: 1.0*50 + 1.4*(-150) + 0.98*(80) = 50 - 210 + 78.4 = -81.6 kN
        # O caso com SCA ZERADA dá: 1.0*50 + 1.4*(-150) + 0.00*(80) = 50 - 210 = -160.0 kN
        # Portanto, V_min deve atingir exatamente -160.0 kN!
        v_min_env = res['envelopes']['V_min']
        self.assertAlmostEqual(v_min_env['value'], -160.0, places=2)
        self.assertIn("SCA Zerada", v_min_env['combination'])

    def test_auto_detect_subsequent_pairs(self):
        """Verifica detecção automática de pares de bases subsequentes (ex: BT1+BT2, BT3+BT4)"""
        bases = ["BT1", "BT2", "BT3", "BT4", "BA1", "BA2"]
        pairs = auto_detect_subsequent_pairs(bases)
        self.assertEqual(len(pairs), 3)
        self.assertEqual(pairs[0]['bases'], ["BA1", "BA2"])
        self.assertEqual(pairs[1]['bases'], ["BT1", "BT2"])
        self.assertEqual(pairs[2]['bases'], ["BT3", "BT4"])

    def test_grouped_bases_combination_and_envelopes(self):
        """Verifica a soma algébrica das forças e a envoltória conjunta para sapatas agrupadas (BT1 + BT2)"""
        sheet_data = {
            'sheet_name': 'Teste_Sapata_Conjunta',
            'bases_data': {
                'BT1': {
                    'PP': {'V': 100.0, 'X': 10.0, 'Y': 5.0},
                    'CVX+': {'V': -40.0, 'X': 25.0, 'Y': 0.0}
                },
                'BT2': {
                    'PP': {'V': 150.0, 'X': -10.0, 'Y': 15.0},
                    'CVX+': {'V': 80.0, 'X': 20.0, 'Y': 0.0}
                }
            }
        }
        options = {
            'base_groups': [
                {'name': 'Sapata 1 (BT1 + BT2)', 'bases': ['BT1', 'BT2']}
            ]
        }

        res = batch_calculate_sheet_bases(sheet_data, standard="NBR 8681", method="ELU_NORMAL", options=options)
        self.assertIn("combined_footings", res)
        cf = res["combined_footings"]
        self.assertEqual(len(cf["summary_table"]), 1)
        row = cf["summary_table"][0]
        self.assertEqual(row["base"], "Sapata 1 (BT1 + BT2)")
        self.assertEqual(row["member_bases"], ["BT1", "BT2"])

        # Para cada combinação, V_total = V_BT1 + V_BT2
        footing_details = cf["footings_details"]["Sapata 1 (BT1 + BT2)"]
        combos = footing_details["combinations"]
        bt1_combos = res["bases_details"]["BT1"]["combinations"]
        bt2_combos = res["bases_details"]["BT2"]["combinations"]

        self.assertEqual(len(combos), len(bt1_combos))
        for i in range(len(combos)):
            expected_v = bt1_combos[i]["V"] + bt2_combos[i]["V"]
            expected_x = bt1_combos[i]["X"] + bt2_combos[i]["X"]
            expected_y = bt1_combos[i]["Y"] + bt2_combos[i]["Y"]
            self.assertAlmostEqual(combos[i]["V"], expected_v, places=2)
            self.assertAlmostEqual(combos[i]["X"], expected_x, places=2)
            self.assertAlmostEqual(combos[i]["Y"], expected_y, places=2)

        # Verificar que a envoltória V_max é a maior soma e contém o breakdown
        v_max_val = row["V_max"]
        self.assertGreater(v_max_val, 0.0)
        self.assertIn("BT1", row["V_max_breakdown"])
        self.assertIn("BT2", row["V_max_breakdown"])
        self.assertAlmostEqual(
            row["V_max_breakdown"]["BT1"]["V"] + row["V_max_breakdown"]["BT2"]["V"],
            v_max_val,
            places=2
        )

    def test_export_excel_with_combined_footings(self):
        """Verifica se a aba 'Sapatas_Conjuntas' é gerada na exportação para Excel"""
        sheet_data = {
            'sheet_name': 'Linha_Transmissao',
            'bases_data': {
                'BT1': {'PP': {'V': 80.0, 'X': 5.0, 'Y': 2.0}},
                'BT2': {'PP': {'V': 90.0, 'X': 4.0, 'Y': 1.0}}
            }
        }
        options = {
            'base_groups': [
                {'name': 'Sapata Perna A (BT1+BT2)', 'bases': ['BT1', 'BT2']}
            ]
        }
        res = batch_calculate_sheet_bases(sheet_data, standard="NBR 8681", method="ELU_NORMAL", options=options)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp_path = tmp.name

        try:
            out_file = export_mietc_results_to_excel(res, tmp_path)
            self.assertTrue(os.path.exists(out_file))

            import openpyxl
            wb = openpyxl.load_workbook(out_file)
            self.assertIn("Sapatas_Conjuntas", wb.sheetnames)
            ws_sc = wb["Sapatas_Conjuntas"]
            self.assertIn("ENVOLTÓRIAS CRÍTICAS DE SAPATAS CONJUNTAS", str(ws_sc.cell(1, 1).value))
            self.assertEqual(ws_sc.cell(6, 1).value, "Sapata Perna A (BT1+BT2)")
            self.assertEqual(ws_sc.cell(6, 2).value, "BT1 + BT2")
            wb.close()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_exclude_wind_from_starting_allow_with_other_actions(self):
        """
        Verifica a regra industrial solicitada:
        - Excluir vento (CVX+, CVY+) e operação (TCO) no momento da partida (TCP).
        - Permitir vento atuando simultaneamente com outras ações (TCO, SCM, SCP, SCA, T).
        - Respeitar coeficientes gamma customizados informados pelo usuário.
        """
        base_loads = {
            'PP': {'V': 100.0, 'X': 0.0, 'Y': 0.0},
            'CP': {'V': 40.0, 'X': 0.0, 'Y': 0.0},
            'SCM': {'V': 25.0, 'X': 0.0, 'Y': 0.0},
            'SCP': {'V': 20.0, 'X': 0.0, 'Y': 0.0},
            'TCO': {'V': 15.0, 'X': 8.0, 'Y': 0.0},
            'TCP': {'V': 10.0, 'X': 12.0, 'Y': 0.0},
            'T': {'V': 0.0, 'X': 5.0, 'Y': 0.0},
            'CVX+': {'V': -30.0, 'X': 35.0, 'Y': 0.0},
            'CVY+': {'V': -25.0, 'X': 0.0, 'Y': 30.0},
            'SCA': {'V': 18.0, 'X': 0.0, 'Y': 0.0},
        }

        options = {
            'action_rules': {
                'PP': {'gamma': 1.35},
                'CP': {'gamma': 1.40},
                'SCM': {'gamma': 1.50},
                'SCP': {'gamma': 1.50},
                'TCO': {'group': 'Operação/Partida', 'excludes': 'TCP', 'gamma': 1.40},
                'TCP': {'group': 'Operação/Partida', 'excludes': 'TCO, CVX+, CVY+', 'gamma': 1.20},
                'T': {'gamma': 1.20},
                'CVX+': {'group': 'Vento', 'excludes': 'TCP', 'gamma': 1.40},
                'CVY+': {'group': 'Vento', 'excludes': 'TCP', 'gamma': 1.40},
                'SCA': {'gamma': 1.50},
            }
        }

        res = calculate_single_base_combinations("BASE_CORREIA", base_loads, standard="NBR 8681", method="ELU_NORMAL", options=options)
        combos = res["combinations"]
        self.assertGreater(len(combos), 0)

        has_tco_and_cvx = False
        has_tco_and_cvy = False
        has_tcp_and_scm = False
        has_cvx_and_scm = False

        for c in combos:
            factors = c.get("factors", {})
            f_tcp = factors.get('TCP', 0.0)
            f_tco = factors.get('TCO', 0.0)
            f_cvx = factors.get('CVX+', 0.0)
            f_cvy = factors.get('CVY+', 0.0)
            f_scm = factors.get('SCM', 0.0)

            # 1. NUNCA pode haver TCP e Vento (CVX+ ou CVY+) na mesma combinação
            if f_tcp > 1e-4:
                self.assertEqual(f_cvx, 0.0, f"Violou regra: TCP ativo com CVX+ na combinação: {c['title']} | {c['formula']}")
                self.assertEqual(f_cvy, 0.0, f"Violou regra: TCP ativo com CVY+ na combinação: {c['title']} | {c['formula']}")
                # 2. NUNCA pode haver TCP e TCO na mesma combinação
                self.assertEqual(f_tco, 0.0, f"Violou regra: TCP ativo com TCO na combinação: {c['title']} | {c['formula']}")
                if f_scm > 1e-4:
                    has_tcp_and_scm = True

            # 3. Vento X e Vento Y nunca simultâneos
            self.assertFalse(f_cvx > 1e-4 and f_cvy > 1e-4, f"Violou XOR Vento: {c['title']}")

            # 4. Vento DEVE poder atuar com TCO
            if f_cvx > 1e-4 and f_tco > 1e-4:
                has_tco_and_cvx = True
            if f_cvy > 1e-4 and f_tco > 1e-4:
                has_tco_and_cvy = True

            # 5. Vento DEVE poder atuar com outras ações (SCM, etc.)
            if f_cvx > 1e-4 and f_scm > 1e-4:
                has_cvx_and_scm = True

            # 6. Coeficientes customizados respeitados quando a ação for principal (fator == gamma)
            if c['type'].startswith('ELU_Normal_Var_TCP_') and f_tcp > 0:
                self.assertAlmostEqual(f_tcp, 1.20, places=2)
            if c['type'].startswith('ELU_Normal_Var_TCO_') and f_tco > 0:
                self.assertAlmostEqual(f_tco, 1.40, places=2)
            if c['type'].startswith('ELU_Normal_Var_SCM_') and f_scm > 0:
                self.assertAlmostEqual(f_scm, 1.50, places=2)

        self.assertTrue(has_tco_and_cvx, "Esperava-se combinações de vento CVX+ com operação TCO")
        self.assertTrue(has_tco_and_cvy, "Esperava-se combinações de vento CVY+ com operação TCO")
        self.assertTrue(has_tcp_and_scm, "Esperava-se combinações de partida TCP com material SCM")
        self.assertTrue(has_cvx_and_scm, "Esperava-se combinações de vento CVX+ com material SCM")

    def test_momento_vertical_and_dynamic_header_detection(self):
        """Verifica a extração de Momento M (kNm), suporte a linha de distância entre extremos e propagação nas combinações e Excel"""
        import io
        import openpyxl

        # 1. Cria planilha simulando exatamente a imagem do usuário:
        # Linha 1: Distância entre extremos | | | BASE 1: 2,5 | BASE 2: 3,0
        # Linha 2: Grupo | Item | Direção | BASE 1 | BASE 2
        # Linha 3: Ações permanentes | PP | V | 2.8 | 3.5
        # Linha 4:                   | PP | M | 12.0 | 15.0
        # Linha 5: Ações permanentes | PP | X | -0.2 | 0.4
        # Linha 6: Ações permanentes | PP | Y | 0.0 | 0.1
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TORRE_1"

        ws.cell(1, 1, "Distância entre extremos")
        ws.cell(1, 4, 2.5)
        ws.cell(1, 5, 3.0)

        ws.cell(2, 1, "Grupo")
        ws.cell(2, 2, "Item")
        ws.cell(2, 3, "Direção")
        ws.cell(2, 4, "BASE 1")
        ws.cell(2, 5, "BASE 2")

        ws.cell(3, 1, "Ações permanentes")
        ws.cell(3, 2, "PP")
        ws.cell(3, 3, "V")
        ws.cell(3, 4, 2.8)
        ws.cell(3, 5, 3.5)

        ws.cell(4, 2, "PP")
        ws.cell(4, 3, "M")
        ws.cell(4, 4, 12.0)
        ws.cell(4, 5, 15.0)

        ws.cell(5, 1, "Ações permanentes")
        ws.cell(5, 2, "PP")
        ws.cell(5, 3, "X")
        ws.cell(5, 4, -0.20)
        ws.cell(5, 5, 0.40)

        ws.cell(6, 1, "Ações permanentes")
        ws.cell(6, 2, "PP")
        ws.cell(6, 3, "Y")
        ws.cell(6, 4, 0.0)
        ws.cell(6, 5, 0.10)

        # Adiciona sobrecarga variável com momento
        ws.cell(7, 1, "Ações variáveis")
        ws.cell(7, 2, "SCM")
        ws.cell(7, 3, "V")
        ws.cell(7, 4, 10.0)
        ws.cell(7, 5, 12.0)

        ws.cell(8, 2, "SCM")
        ws.cell(8, 3, "M")
        ws.cell(8, 4, 5.0)
        ws.cell(8, 5, 6.0)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        # Teste parsing
        parsed = parse_mietc_workbook(stream)
        self.assertIn("TORRE_1", parsed["sheets"])
        sheet = parsed["sheets"]["TORRE_1"]
        self.assertEqual(sheet["bases_count"], 2)
        self.assertIn("BASE 1", sheet["bases_list"])
        self.assertIn("BASE 2", sheet["bases_list"])

        b1_loads = sheet["bases_data"]["BASE 1"]
        self.assertIn("PP", b1_loads)
        self.assertEqual(b1_loads["PP"]["V"], 2.8)
        self.assertEqual(b1_loads["PP"]["M"], 12.0)
        self.assertEqual(b1_loads["PP"]["X"], -0.20)
        self.assertEqual(b1_loads["PP"]["Y"], 0.0)

        self.assertIn("SCM", b1_loads)
        self.assertEqual(b1_loads["SCM"]["M"], 5.0)

        # Teste cálculo de combinações com M
        opts = {
            'base_groups': [
                {'id': 'S1', 'name': 'Sapata 1', 'bases': ['BASE 1', 'BASE 2']}
            ]
        }
        res = batch_calculate_sheet_bases(sheet, standard="NBR 8681", method="ELU_NORMAL", options=opts)
        self.assertEqual(res["total_bases"], 2)

        # Verifica combinações da BASE 1
        b1_details = res["bases_details"]["BASE 1"]
        combos = b1_details["combinations"]
        self.assertGreater(len(combos), 0)

        for c in combos:
            self.assertIn("M", c)
            # Para combinações de pura permanente: 1.35 * PP (2.8 V, 12.0 M)
            factors = c.get("factors", {})
            f_pp = factors.get("PP", 0.0)
            f_scm = factors.get("SCM", 0.0)
            expected_m = round(f_pp * 12.0 + f_scm * 5.0, 2)
            self.assertAlmostEqual(c["M"], expected_m, places=2)

        # Verifica envoltórias da BASE 1
        env1 = b1_details["envelopes"]
        self.assertIn("M_max", env1["peaks"])
        self.assertGreater(env1["peaks"]["M_max"], 0.0)
        self.assertIn("concomitant_M", env1["V_max"])
        self.assertIn("concomitant_M", env1["V_min"])
        self.assertIn("concomitant_M", env1["X_max"])
        self.assertIn("concomitant_M", env1["Y_max"])
        self.assertIn("concomitant_M", env1["H_max"])

        # Verifica envoltórias de sapatas conjuntas (soma simultânea de M)
        cf = res["combined_footings"]
        cf_summary = cf["summary_table"]
        self.assertEqual(len(cf_summary), 1)
        s1 = cf_summary[0]
        self.assertIn("peak_M", s1)
        self.assertGreater(s1["peak_M"], 0.0)
        self.assertIn("V_max_M", s1)
        self.assertIn("V_min_M", s1)

        # Verifica exportação Excel com abas e colunas de Momento M
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        out_path = export_mietc_results_to_excel(res, tmp_path)
        wb_exp = openpyxl.load_workbook(out_path)

        # 1. ws_cargas tem M_k
        self.assertIn("Resumo_Cargas_Consideradas", wb_exp.sheetnames)
        ws_cargas = wb_exp["Resumo_Cargas_Consideradas"]
        # Procurar por "M_k (kNm)"
        found_mk = any("M_k (kNm)" in str(cell.value) for row in ws_cargas.iter_rows() for cell in row)
        self.assertTrue(found_mk, "Coluna M_k (kNm) não encontrada em Resumo_Cargas_Consideradas")

        # 2. ws_sum tem Momento M
        self.assertIn("Resumo_Envoltórias", wb_exp.sheetnames)
        ws_sum = wb_exp["Resumo_Envoltórias"]
        found_m_peak = any("1. |M| Máx Pico (kNm)" in str(cell.value) for row in ws_sum.iter_rows() for cell in row)
        self.assertTrue(found_m_peak, "1. |M| Máx Pico (kNm) não encontrado em Resumo_Envoltórias")

        # 3. ws_comb tem Momento M
        self.assertIn("Sapatas_Conjuntas", wb_exp.sheetnames)
        ws_comb = wb_exp["Sapatas_Conjuntas"]
        found_comb_m = any("1. |M| Máx Pico (kNm)" in str(cell.value) for row in ws_comb.iter_rows() for cell in row)
        self.assertTrue(found_comb_m, "1. |M| Máx Pico (kNm) não encontrado em Sapatas_Conjuntas")

        # 4. ws_all tem Momento M (kNm)
        self.assertIn("Todas_Combinações", wb_exp.sheetnames)
        ws_all = wb_exp["Todas_Combinações"]
        self.assertEqual(ws_all.cell(1, 6).value, "Momento M (kNm)")

        wb_exp.close()
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    def test_similarity_clustering_and_excel_tab(self):
        """Verifica a tipificação por similaridade de bases e a geração da aba Similaridade_Dimensionamento no Excel."""
        import openpyxl
        import tempfile
        import os
        from backend.calculators.comb_bases_mietc import cluster_bases_by_similarity, export_mietc_results_to_excel, batch_calculate_sheet_bases

        # 1. Teste direto do agrupador de similaridade
        mock_summary_items = [
            {"base": "B1", "V_max": 100.0, "V_min": 20.0, "H_max": 25.0, "M_max": 10.0, "comb_V_max": "1.4PP", "comb_V_min": "1.0PP", "comb_H_max": "1.4W"},
            {"base": "B2", "V_max": 95.0, "V_min": 18.0, "H_max": 23.0, "M_max": 9.0, "comb_V_max": "1.4PP", "comb_V_min": "1.0PP", "comb_H_max": "1.4W"}, # similar a B1 (dentro de 20%)
            {"base": "B3", "V_max": 50.0, "V_min": 5.0, "H_max": 12.0, "M_max": 4.0, "comb_V_max": "1.4PP", "comb_V_min": "1.0PP", "comb_H_max": "1.4W"}, # carga menor -> outro grupo
            {"base": "B_TRAÇÃO_1", "V_max": 40.0, "V_min": -35.0, "H_max": 30.0, "M_max": 15.0, "comb_V_max": "1.4PP", "comb_V_min": "1.0PP+1.4W", "comb_H_max": "1.4W"}, # tração
            {"base": "B_TRAÇÃO_2", "V_max": 38.0, "V_min": -32.0, "H_max": 28.0, "M_max": 14.0, "comb_V_max": "1.4PP", "comb_V_min": "1.0PP+1.4W", "comb_H_max": "1.4W"}, # similar tração
        ]

        res_clustering = cluster_bases_by_similarity(mock_summary_items, tolerance=0.20)
        clusters = res_clustering["families"]
        self.assertGreaterEqual(len(clusters), 3)

        # Deve haver famílias de tração (prefixo T) e compressão (prefixo S)
        tension_fams = [c for c in clusters if c["mode"] == "TRAÇÃO"]
        comp_fams = [c for c in clusters if c["mode"] == "COMPRESSÃO"]
        self.assertTrue(len(tension_fams) >= 1)
        self.assertTrue(len(comp_fams) >= 2)

        # B1 e B2 devem estar na mesma família de compressão, sendo B1 a governante
        fam_b1 = next(c for c in comp_fams if "B1" in c["bases"])
        self.assertIn("B2", fam_b1["bases"])
        self.assertEqual(fam_b1["governing_base"], "B1")
        self.assertEqual(fam_b1["governing_loads"]["V_max"], 100.0)

        # B_TRAÇÃO_1 e B_TRAÇÃO_2 devem estar na mesma família de tração
        fam_t1 = next(c for c in tension_fams if "B_TRAÇÃO_1" in c["bases"])
        self.assertIn("B_TRAÇÃO_2", fam_t1["bases"])
        self.assertEqual(fam_t1["governing_base"], "B_TRAÇÃO_1")

        # 2. Teste da geração da aba de Similaridade no Excel
        fake_sheet = {
            "sheet_name": "Folha_Teste",
            "bases_data": {
                "B1": {"PP": {"V": 70.0, "X": 5.0, "Y": 5.0, "M": 0.0}},
                "B2": {"PP": {"V": 65.0, "X": 4.5, "Y": 4.5, "M": 0.0}},
                "B3": {"PP": {"V": 20.0, "X": 2.0, "Y": 2.0, "M": 0.0}},
            }
        }
        batch_res = batch_calculate_sheet_bases(fake_sheet, standard="NBR 8681", method="ELU_NORMAL")
        self.assertIn("similarity_clustering", batch_res)
        self.assertGreater(batch_res["similarity_clustering"]["total_families"], 0)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        out_path = export_mietc_results_to_excel(batch_res, tmp_path)
        wb_exp = openpyxl.load_workbook(out_path)

        self.assertIn("Similaridade_Dimensionamento", wb_exp.sheetnames)
        ws_sim = wb_exp["Similaridade_Dimensionamento"]

        # Verifica cabeçalhos e estrutura das duas tabelas
        cell_vals = [str(cell.value) for row in ws_sim.iter_rows() for cell in row if cell.value is not None]
        self.assertTrue(any("SÍNTESE DAS FAMÍLIAS TÍPICAS" in v for v in cell_vals))
        self.assertTrue(any("RELAÇÃO NOMINAL DE TODAS AS BASES" in v for v in cell_vals))
        self.assertTrue(any("% em Relação ao Tipo" in v for v in cell_vals))

        # Verifica presença de bordas nas células de dados
        has_borders = any(
            cell.border is not None and cell.border.top.style is not None
            for row in ws_sim.iter_rows(min_row=5, max_row=15)
            for cell in row
        )
        self.assertTrue(has_borders, "Células da planilha Similaridade_Dimensionamento devem conter bordas")

        wb_exp.close()
        try:
            os.remove(tmp_path)
        except OSError:
            pass

if __name__ == "__main__":
    unittest.main()

