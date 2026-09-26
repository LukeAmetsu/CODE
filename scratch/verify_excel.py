import openpyxl

wb = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\Refs\Empuxos_Atualizados_Interfaces.xlsx', data_only=True)
ws = wb['Empuxos_Separados_STRAP']

tot_Fa = sum(ws.cell(r, 25).value for r in range(37, 66))
tot_Fp = sum(ws.cell(r, 26).value for r in range(37, 66))
M_tom = sum(ws.cell(r, 25).value * (12.0 - ws.cell(r, 1).value) for r in range(37, 66))
M_res = sum(ws.cell(r, 26).value * (12.0 - ws.cell(r, 1).value) for r in range(37, 66))

u27 = ws['U27'].value
u28 = ws['U28'].value
u29 = ws['U29'].value
u30 = ws['U30'].value
u31 = ws['U31'].value
u32 = ws['U32'].value

print('Sum Fa:', tot_Fa, 'vs U27:', u27)
print('Sum Fp:', tot_Fp, 'vs U28:', u28)
print('M_tom:', M_tom, 'vs U29:', u29)
print('M_res:', M_res, 'vs U30:', u30)
print('FS_desl:', tot_Fp/tot_Fa, 'vs U31:', u31)
print('FS_tomb:', M_res/M_tom, 'vs U32:', u32)
