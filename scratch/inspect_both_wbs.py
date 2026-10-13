import openpyxl

def dump_wb(path, outfile):
    wb = openpyxl.load_workbook(path, data_only=True)
    with open(outfile, 'w', encoding='utf-8') as f:
        f.write(f"=== WORKBOOK: {path} ===\n")
        f.write(f"Sheets: {wb.sheetnames}\n\n")
        for name in wb.sheetnames:
            ws = wb[name]
            f.write(f"--- Sheet [{name}] (rows={ws.max_row}, cols={ws.max_column}) ---\n")
            for r in range(1, min(15, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    f.write(f"  Row {r:2d}: {row_vals}\n")
            f.write("\n")

dump_wb(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\Results_FS_2.5-2.75.xlsx', 'scratch/dump_results_fs.txt')
dump_wb(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\capacidade_estaca_export.xlsx', 'scratch/dump_export.txt')
print("Dumped both workbooks to scratch/!")
