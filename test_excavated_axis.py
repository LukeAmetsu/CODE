import openpyxl
from backend.calculators.capacidade_estaca import (
    calculate_layer_stresses_and_strengths,
    DEFAULT_SOIL_GROUPS
)

wb1 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\Results_FS_2.5-2.75.xlsx', data_only=True)
s1 = wb1['Summary axial lengths']
wb1_data = {}
for r in list(s1.iter_rows(values_only=True))[1:]:
    ax = str(r[0]).strip()
    L = r[4]
    if L in [5, 10, 15, 20, 25, 30]:
        wb1_data[(ax, L)] = {
            'qs_adm': r[12],
            'qp_adm': r[13],
            'w': r[15],
            'qnet': r[16],
            'fs_avg': r[8]
        }

wb2 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\capacidade_estaca_export.xlsx', data_only=True)
s_spt = wb2['SPT-01']
# Load raw SPT-01 layers
raw_layers = []
for r in list(s_spt.iter_rows(values_only=True))[2:]:
    if r[0] is not None and r[1] is not None:
        raw_layers.append({
            'z_ini': float(r[0]),
            'z_fim': float(r[1]),
            'n1_60': float(r[2]),
            'soil_group': str(r[3]),
            'sucs': str(r[4])
        })

print(f'Loaded {len(raw_layers)} raw layers from SPT-01')

# For Axis A: cota_topo = 20.40m
# If soil above cota_topo is excavated:
cota_topo = 20.40
excavated_layers = []
for l in raw_layers:
    if l['z_fim'] <= cota_topo:
        continue
    # Trim or shift
    z_ini_new = max(0.0, l['z_ini'] - cota_topo)
    z_fim_new = l['z_fim'] - cota_topo
    excavated_layers.append({
        'z_ini': z_ini_new,
        'z_fim': z_fim_new,
        'n1_60': l['n1_60'],
        'soil_group': l['soil_group'],
        'sucs': l['sucs']
    })

print(f'Excavated layers for Axis A: {len(excavated_layers)} layers, starting at z_ini={excavated_layers[0]["z_ini"]}')

# Let's test with AASHTO 2007 method and Mayne method
for lat_method in ['aashto_2007', 'mayne_fhwa']:
    proc = calculate_layer_stresses_and_strengths(
        excavated_layers,
        water_table_m=0.0, # All submerged below seabed
        soil_groups=DEFAULT_SOIL_GROUPS,
        lateral_method=lat_method,
        tip_method='aashto_2007',
        apply_qs_cap=True,
        apply_qp_cap=True
    )
    
    # Calculate cumulative Qs for D=1.2m with 4m casing (-40%)
    perim = 3.14159265 * 1.2
    casing_len = 4.0
    casing_factor = 0.60
    
    print(f'\n--- Method: {lat_method} ---')
    print(f'{"L(m)":<6} | {"fs_mid (kPa)":<14} | {"Qs_adm (kN)":<14} | {"WB1 Qs_adm":<14} | {"Diff(%)":<10}')
    
    cum_qs = 0.0
    for idx, lyr in enumerate(proc):
        dz = lyr['dz']
        z_mid = lyr['z_mid']
        qs_unit = lyr['qs']
        
        # Apply casing factor
        if z_mid <= casing_len:
            eff_qs = qs_unit * casing_factor
        else:
            eff_qs = qs_unit
            
        cum_qs += eff_qs * perim * dz
        L = round(lyr['z_fim'], 1)
        if L in [5.0, 10.0, 15.0, 20.0, 25.0, 30.0]:
            qs_adm = cum_qs / 2.5
            wb1_val = wb1_data.get(('A', int(L)), {}).get('qs_adm', 0.0)
            diff = (qs_adm - wb1_val) / wb1_val * 100 if wb1_val else 0
            print(f'{L:<6.1f} | {qs_unit:<14.2f} | {qs_adm:<14.1f} | {wb1_val:<14.1f} | {diff:<+9.1f}%')
