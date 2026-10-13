import openpyxl
from backend.calculators.capacidade_estaca import (
    calculate_layer_stresses_and_strengths,
    DEFAULT_SOIL_GROUPS
)

wb1 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\Results_FS_2.5-2.75.xlsx', data_only=True)
s_wb1 = wb1['Axis A axial lengths']
wb1_dict = {}
for r in list(s_wb1.iter_rows(values_only=True))[1:]:
    wb1_dict[r[0]] = {'qs': r[2], 'fs': r[3], 'qp': r[5], 'qall_s': r[8], 'w': r[11], 'qnet': r[12]}

wb2 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\capacidade_estaca_export.xlsx', data_only=True)
s_spt = wb2['SPT-01']
raw_layers = []
for r in list(s_spt.iter_rows(values_only=True))[2:]:
    if r[0] is not None and r[1] is not None:
        raw_layers.append({
            'z_ini': float(r[0]),
            'z_fim': float(r[1]),
            'n1_60': float(r[2]),
            'soil_group': str(r[3]),
            'sucs': str(r[4])
        })

cota_topo = 20.40
axis_layers = []
for l in raw_layers:
    if l['z_fim'] <= cota_topo:
        continue
    z_a = max(l['z_ini'], cota_topo) - cota_topo
    z_b = l['z_fim'] - cota_topo
    axis_layers.append({
        'z_ini': z_a,
        'z_fim': z_b,
        'n1_60': l['n1_60'],
        'soil_group': l['soil_group'],
        'sucs': l['sucs']
    })

with open('scratch_test_sliced.txt', 'w', encoding='utf-8') as f:
    f.write(f'Sliced {len(axis_layers)} layers starting from cota_topo = 20.40m\n\n')
    for m in ['mayne_fhwa', 'aashto_2007']:
        proc = calculate_layer_stresses_and_strengths(
            axis_layers,
            water_table_m=0.0,
            lateral_method=m,
            tip_method='aashto_2007',
            apply_qs_cap=True,
            apply_qp_cap=True
        )
        f.write(f'=== Method: {m} ===\n')
        f.write(f'{"L(m)":<6} | {"fs_calc":<10} {"fs_WB1":<10} | {"qp_calc":<10} {"qp_WB1":<10} | {"sigma_v":<10}\n')
        f.write('-'*60 + '\n')
        for L in [1, 2, 3, 4, 5, 6, 8, 10, 15, 20]:
            matching = [lyr for lyr in proc if abs(lyr['z_fim'] - L) < 0.35]
            if matching:
                lyr = matching[0]
                wb1_val = wb1_dict.get(L, {})
                fs_w = wb1_val.get('fs', 0.0)
                qp_w = wb1_val.get('qp', 0.0)
                f.write(f'{L:<6.1f} | {lyr["qs"]:<10.2f} {fs_w:<10.2f} | {lyr["qp"]:<10.2f} {qp_w:<10.2f} | {lyr["sigma_v"]:<10.2f}\n')
        f.write('\n')

print('Test completed successfully.')
