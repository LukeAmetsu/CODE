import openpyxl
from backend.calculators.capacidade_estaca import (
    calculate_layer_stresses_and_strengths,
    DEFAULT_SOIL_GROUPS
)

wb1 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\Results_FS_2.5-2.75.xlsx', data_only=True)
s_wb1 = wb1['Summary axial lengths']
wb1_data = {}
for r in list(s_wb1.iter_rows(values_only=True))[1:]:
    ax = str(r[0]).strip()
    L = r[4]
    wb1_data[(ax, L)] = {
        'qs_adm': r[12],
        'qp_adm': r[13],
        'w': r[15],
        'qnet': r[16],
        'fs_avg': r[8]
    }

wb2 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\capacidade_estaca_export.xlsx', data_only=True)
s_spt = wb2['SPT-01']
raw_layers = []
for r in list(s_spt.iter_rows(values_only=True))[2:]:
    if r[0] is not None and r[1] is not None:
        raw_layers.append({
            'z_ini': float(r[0]),
            'z_fim': float(r[1]),
            'n1_60': float(r[2]),
            'soil_group': str(r[3]),
            'sucs': str(r[4])
        })

axes_info = {
    'A': {'cota_topo': 20.40, 'D': 1.2, 'L_above': 22.92},
    'B': {'cota_topo': 18.60, 'D': 1.2, 'L_above': 21.12},
    'C': {'cota_topo': 13.82, 'D': 1.2, 'L_above': 16.34},
    'D': {'cota_topo': 8.94,  'D': 1.2, 'L_above': 11.46},
    'E': {'cota_topo': 4.16,  'D': 1.5, 'L_above': 6.68},
}

with open('scratch_all_axes_excavated.txt', 'w', encoding='utf-8') as f:
    f.write('=== SIMULAÇÃO COM SOLO ACIMA DA COTA DE CORTE REMOVIDO ===\n\n')
    
    for ax_name, ax_cfg in axes_info.items():
        cota_topo = ax_cfg['cota_topo']
        D = ax_cfg['D']
        perim = 3.14159265 * D
        area = 3.14159265 * (D**2) / 4.0
        
        # Slice layers
        axis_layers = []
        for l in raw_layers:
            if l['z_fim'] <= cota_topo:
                continue
            z_a = max(l['z_ini'], cota_topo) - cota_topo
            z_b = l['z_fim'] - cota_topo
            axis_layers.append({
                'z_ini': z_a,
                'z_fim': z_b,
                'n1_60': l['n1_60'],
                'soil_group': l['soil_group'],
                'sucs': l['sucs']
            })
            
        proc = calculate_layer_stresses_and_strengths(
            axis_layers,
            water_table_m=0.0,
            lateral_method='aashto_2007',
            tip_method='aashto_2007',
            apply_qs_cap=True,
            apply_qp_cap=True
        )
        
        f.write(f'================ EIXO {ax_name} (Topo: {cota_topo}m, D={D}m) ================\n')
        f.write(f'{"L(m)":<6} | {"Qs_adm Prog":<12} {"Qs_adm WB1":<12} | {"Qp_adm Prog":<12} {"Qp_adm WB1":<12} | {"Qnet Prog":<12} {"Qnet WB1":<12} | {"Diff Qnet(%)":<12}\n')
        f.write('-'*95 + '\n')
        
        # Track cumulative Qs
        cum_qs = 0.0
        casing_len = 4.0 # Designer uses 4m
        casing_factor = 0.60 # -40%
        
        for lyr in proc:
            dz = lyr['dz']
            z_mid = lyr['z_mid']
            qs_unit = lyr['qs']
            if z_mid <= casing_len:
                eff_qs = qs_unit * casing_factor
            else:
                eff_qs = qs_unit
                
            cum_qs += eff_qs * perim * dz
            L = round(lyr['z_fim'], 1)
            
            if L in [2.0, 5.0, 10.0, 15.0, 20.0, 25.0, 30.0, 35.0]:
                qs_adm = cum_qs / 2.5
                qp_unit = lyr['qp']
                qp_adm = (qp_unit * area) / 2.75
                # Pile weight: submerged concrete gamma = 14 kN/m3
                w_pile = area * (L + ax_cfg['L_above']) * 14.0
                qnet = qs_adm + qp_adm - w_pile
                
                wb1_entry = wb1_data.get((ax_name, int(L)), {})
                wb1_qs = wb1_entry.get('qs_adm', 0.0)
                wb1_qp = wb1_entry.get('qp_adm', 0.0)
                wb1_qnet = wb1_entry.get('qnet', 0.0)
                
                diff = (qnet - wb1_qnet) / abs(wb1_qnet) * 100 if wb1_qnet else 0
                f.write(f'{L:<6.1f} | {qs_adm:<12.1f} {wb1_qs:<12.1f} | {qp_adm:<12.1f} {wb1_qp:<12.1f} | {qnet:<12.1f} {wb1_qnet:<12.1f} | {diff:<+11.1f}%\n')
        f.write('\n')

print('All axes simulation finished.')
