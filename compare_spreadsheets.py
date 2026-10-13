import openpyxl

wb1 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\Results_FS_2.5-2.75.xlsx', data_only=True)
wb2 = openpyxl.load_workbook(r'C:\OD\OneDrive - Andrade Gutierrez\CODE-2\dist\capacidade_estaca_export.xlsx', data_only=True)

with open('scratch_compare_results.txt', 'w', encoding='utf-8') as f:
    f.write('=== DETAILED COMPARISON: Results_FS_2.5-2.75.xlsx vs capacidade_estaca_export.xlsx ===\n\n')
    
    # 1. Compare Settings / Assumptions
    f.write('1. SETTINGS & PREMISSAS:\n')
    if 'Settings' in wb1.sheetnames:
        for r in wb1['Settings'].iter_rows(values_only=True):
            if any(r):
                f.write(f'  [Projetista WB1] {r[0]}: {r[1]}\n')
    f.write('\n')

    s1 = wb1['Summary axial lengths']
    rows1 = list(s1.iter_rows(values_only=True))
    h1 = rows1[0]
    
    wb1_axis_data = {}
    for r in rows1[1:]:
        axis = str(r[0]).strip()
        if axis not in wb1_axis_data:
            wb1_axis_data[axis] = []
        wb1_axis_data[axis].append(r)
        
    for ax in ['A', 'B', 'C', 'D', 'E']:
        f.write(f'================ AXIS {ax} ================\n')
        d1 = wb1_axis_data.get(ax, [])
        sheet_name = f'Axis {ax}'
        
        if sheet_name in wb2.sheetnames:
            s2 = wb2[sheet_name]
            rows2 = list(s2.iter_rows(values_only=True))
            meta2 = rows2[1][0] if len(rows2) > 1 else ''
            f.write(f'Export Program (WB2) Metadata: {meta2}\n\n')
            
            # Map WB2 by embedment depth
            wb2_map = {}
            for r in rows2[3:]:
                if r[1] is not None:
                    try:
                        emb = round(float(r[1]), 2)
                        wb2_map[emb] = r
                    except:
                        pass
                        
            # Determine diameter index in WB2
            # Check what diameter WB1 uses for this axis
            d_val_wb1 = d1[0][1] if d1 else 1.2
            if abs(d_val_wb1 - 1.5) < 0.05:
                d_idx = (18, 19, 20, 21, 25) # 1.50m
                d_str = "D=1.50m"
            elif abs(d_val_wb1 - 1.3) < 0.05:
                d_idx = (10, 11, 12, 13, 17) # 1.30m
                d_str = "D=1.30m"
            else:
                d_idx = (2, 3, 4, 5, 9)       # 1.20m
                d_str = "D=1.20m"
                
            f.write(f'Diameter in WB1: {d_val_wb1}m ({d_str})\n')
            f.write(f'{"L(m)":<6} | {"Rs_adm WB1":<12} {"Rs_adm WB2":<12} | {"Rp_adm WB1":<12} {"Rp_adm WB2":<12} | {"W_pile WB1":<12} {"W_pile WB2":<12} | {"Qnet WB1":<12} {"Qnet WB2":<12} | {"Diff(%)":<10}\n')
            f.write('-'*105 + '\n')
            
            for r1 in d1:
                L = r1[4]
                if L is not None and L in [0, 2, 5, 10, 15, 20, 25, 30, 35, 37]:
                    qs1 = r1[6]
                    qp1 = r1[10]
                    qall_s1 = r1[12] if r1[12] is not None else 0
                    qall_p1 = r1[13] if r1[13] is not None else 0
                    qgross1 = r1[14] if r1[14] is not None else 0
                    w1 = r1[15] if r1[15] is not None else 0
                    qnet1 = r1[16] if r1[16] is not None else 0
                    
                    matched_r2 = None
                    for emb_k, val in wb2_map.items():
                        if abs(emb_k - float(L)) < 0.35:
                            matched_r2 = val
                            break
                            
                    if matched_r2:
                        qnet2 = float(matched_r2[d_idx[0]])
                        rs2 = float(matched_r2[d_idx[1]])
                        rp2 = float(matched_r2[d_idx[2]])
                        qgross2 = float(matched_r2[d_idx[3]])
                        w2 = float(matched_r2[d_idx[4]])
                        
                        diff_pct = ((qnet2 - qnet1) / abs(qnet1) * 100) if qnet1 != 0 else 0
                        f.write(f'{float(L):<6.1f} | {qall_s1:<12.1f} {rs2:<12.1f} | {qall_p1:<12.1f} {rp2:<12.1f} | {w1:<12.1f} {w2:<12.1f} | {qnet1:<12.1f} {qnet2:<12.1f} | {diff_pct:<+9.2f}%\n')
                    else:
                        f.write(f'{float(L):<6.1f} | {qall_s1:<12.1f} {"N/A":<12} | {qall_p1:<12.1f} {"N/A":<12} | {w1:<12.1f} {"N/A":<12} | {qnet1:<12.1f} {"N/A":<12} | N/A\n')
            f.write('\n')

print('Comparison finished successfully.')
